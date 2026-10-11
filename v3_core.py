"""
magazine_core.py - Styling, tokens, and reusable components for CE Annual Magazine 2025-26
"""

import os, sys, io
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ── Color Palette ──
NAVY_HEX = '1B3A5C'
NAVY_DARK_HEX = '0F2440'
GOLD_HEX = 'C8A962'
GOLD_LIGHT_HEX = 'E8DCC0'
GOLD_BG_HEX = 'FBF7EE'
GREEN_HEX = '2E7D32'
WHITE_HEX = 'FFFFFF'
CREAM_HEX = 'FAFAF8'
LIGHT_BG = 'F5F3EE'
ALT_ROW = 'F3EFE6'
BORDER_HEX = 'D8D4CC'
MUTED_HEX = '6B7280'

NAVY = RGBColor(0x1B, 0x3A, 0x5C)
NAVY_DARK = RGBColor(0x0F, 0x24, 0x40)
GOLD = RGBColor(0xC8, 0xA9, 0x62)
GOLD_DARK = RGBColor(0xA0, 0x88, 0x40)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x4A, 0x55, 0x68)
LIGHT_GREY = RGBColor(0x71, 0x80, 0x96)

def create_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    # Base style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)
    style.font.color.rgb = BLACK
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15

    for level, (sz, sp_b) in {1: (18, 0), 2: (13, 10), 3: (11, 6)}.items():
        h = doc.styles[f'Heading {level}']
        h.font.name = 'Cambria'
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = NAVY
        h.paragraph_format.space_before = Pt(sp_b)
        h.paragraph_format.space_after = Pt(4)

    return doc

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    b_xml = [f'<w:tcBorders {nsdecls("w")}>']
    for side, border in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if border:
            c = border.get('color', 'auto')
            sz = border.get('sz', '4')
            val = border.get('style', 'single')
            b_xml.append(f'<w:{side} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{c}"/>')
        else:
            b_xml.append(f'<w:{side} w:val="none"/>')
    b_xml.append('</w:tcBorders>')
    tcPr.append(parse_xml(''.join(b_xml)))

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def make_table_full_width(table):
    tblPr = table._tbl.tblPr
    tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:w="5000" w:type="pct"/>')
    tblPr.append(tblW)

def run_styled(paragraph, text, font_name='Calibri', size=10, bold=False, italic=False, color=BLACK):
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return run

def add_spacer(doc, pts=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(pts)
    p.paragraph_format.line_spacing = 1.0

def section_header(doc, title, subtitle=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_table_full_width(t)
    cell = t.cell(0, 0)
    set_cell_shading(cell, NAVY_HEX)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_borders(cell,
        bottom={'color': GOLD_HEX, 'sz': '18', 'style': 'single'},
        left={'color': GOLD_HEX, 'sz': '24', 'style': 'single'})

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    if subtitle:
        run_styled(p, subtitle.upper() + '\n', 'Calibri', 8.5, bold=True, color=GOLD)
    run_styled(p, title, 'Cambria', 17, bold=True, color=WHITE)
    add_spacer(doc, 8)

def sub_section_header(doc, title, num=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    if num:
        run_styled(p, f"{num}  ", 'Calibri', 11, bold=True, color=GOLD_DARK)
    run_styled(p, title, 'Cambria', 12.5, bold=True, color=NAVY)

def stat_cards(doc, stats):
    n = len(stats)
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_table_full_width(t)
    for j, (number, label) in enumerate(stats):
        cell = t.cell(0, j)
        set_cell_shading(cell, LIGHT_BG)
        set_cell_margins(cell, top=100, bottom=100, left=60, right=60)
        set_cell_borders(cell,
            bottom={'color': NAVY_HEX, 'sz': '12', 'style': 'single'},
            top={'color': BORDER_HEX, 'sz': '4', 'style': 'single'},
            left={'color': BORDER_HEX, 'sz': '4', 'style': 'single'},
            right={'color': BORDER_HEX, 'sz': '4', 'style': 'single'})
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(2)
        run_styled(p1, str(number), 'Cambria', 17, bold=True, color=NAVY)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        run_styled(p2, label.upper(), 'Calibri', 7.5, bold=True, color=GREY)
    add_spacer(doc, 8)

def callout_box(doc, text, title=None, border_color=NAVY_HEX, bg_color=LIGHT_BG):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_table_full_width(t)
    cell = t.cell(0, 0)
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_cell_borders(cell, left={'color': border_color, 'sz': '24', 'style': 'single'})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    if title:
        run_styled(p, title.upper() + '\n', 'Calibri', 8.5, bold=True, color=NAVY)
    run_styled(p, text, 'Calibri', 9.5, italic=True, color=BLACK)
    add_spacer(doc, 6)

def data_table(doc, headers, rows, col_widths=None, compact=False):
    sz = 8.5 if compact else 9.5
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_table_full_width(t)

    # Repeat header row on every page & prevent row splits across pages
    header_tr = t.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    # Header Row
    for j, h in enumerate(headers):
        cell = t.cell(0, j)
        set_cell_shading(cell, NAVY_HEX)
        pad = 55 if compact else 75
        set_cell_margins(cell, top=pad, bottom=pad, left=60, right=60)
        set_cell_borders(cell,
            bottom={'color': GOLD_HEX, 'sz': '12', 'style': 'single'},
            top={'color': NAVY_HEX, 'sz': '4', 'style': 'single'})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_styled(p, str(h).upper(), 'Calibri', sz, bold=True, color=WHITE)

    # Data Rows
    for i, row in enumerate(rows):
        bg = ALT_ROW if i % 2 == 1 else WHITE_HEX
        for j, val in enumerate(row):
            cell = t.cell(i + 1, j)
            set_cell_shading(cell, bg)
            pad = 40 if compact else 50
            set_cell_margins(cell, top=pad, bottom=pad, left=60, right=60)
            set_cell_borders(cell,
                bottom={'color': BORDER_HEX, 'sz': '2', 'style': 'single'},
                left={'color': BORDER_HEX, 'sz': '1', 'style': 'single'},
                right={'color': BORDER_HEX, 'sz': '1', 'style': 'single'})
            p = cell.paragraphs[0]
            bold = (j == 0 and len(headers) > 3)
            cell_val = str(val).strip() if val is not None else ""
            if not cell_val:
                cell_val = "-"
            run_styled(p, cell_val, 'Calibri', sz, bold=bold, color=BLACK)

    if col_widths:
        for row in t.rows:
            for j, w in enumerate(col_widths):
                if j < len(row.cells):
                    row.cells[j].width = Inches(w)

    add_spacer(doc, 6)

def body(doc, text, bold=False, italic=False, color=BLACK, indent=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if indent:
        p.paragraph_format.left_indent = Inches(0.2)
    run_styled(p, text, 'Calibri', 10, bold=bold, italic=italic, color=color)
    return p

def bullet(doc, text, bold_prefix=None, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    run_styled(p, "• ", 'Calibri', 10, bold=True, color=NAVY)
    if bold_prefix:
        run_styled(p, bold_prefix + " ", 'Calibri', 10, bold=True, color=BLACK)
    run_styled(p, text, 'Calibri', 10, color=BLACK)
    return p

def add_image_safe(doc, path, width=Inches(5.5), caption_text=None, alignment=WD_ALIGN_PARAGRAPH.CENTER):
    if not os.path.exists(path):
        return None
    try:
        p = doc.add_paragraph()
        p.alignment = alignment
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(path, width=width)
        if caption_text:
            cp = doc.add_paragraph()
            cp.alignment = alignment
            cp.paragraph_format.space_before = Pt(1)
            cp.paragraph_format.space_after = Pt(6)
            run_styled(cp, caption_text, 'Calibri', 8.5, italic=True, color=GREY)
        return p
    except Exception as e:
        print(f"Warning: Failed to add image {path}: {e}")
        return None
