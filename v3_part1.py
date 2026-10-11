"""
verified_builder_part1.py - Sections 1 to 4: Leadership, Vision, Mission, PEOs, POs, PSOs & Audited Attainment
"""

import os
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from magazine_core import (
    NAVY_HEX, NAVY_DARK_HEX, GOLD_HEX, GOLD_LIGHT_HEX, GOLD_BG_HEX,
    WHITE_HEX, LIGHT_BG, ALT_ROW, BORDER_HEX, NAVY, GOLD, GOLD_DARK,
    WHITE, BLACK, GREY, LIGHT_GREY,
    run_styled, add_spacer, section_header, sub_section_header,
    stat_cards, callout_box, data_table, body, bullet, add_image_safe
)

BASE = r'C:\Users\Nirupam\Documents\magazine ce'
IMG_DIR = os.path.join(BASE, 'extracted_images')
MEDIA_DIR = r'C:\Users\Nirupam\.gemini\antigravity-ide\brain\d84a1836-479b-4390-be26-b801a6bee109\scratch\baseline_media'

def build_cover_and_frontmatter(doc):
    # Top Spacer
    add_spacer(doc, 36)

    # Institution Header
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    run_styled(p_inst, "BHARATIYA VIDYA BHAVAN'S\n", 'Calibri', 13, bold=True, color=GOLD_DARK)
    run_styled(p_inst, "SARDAR PATEL INSTITUTE OF TECHNOLOGY\n", 'Cambria', 20, bold=True, color=NAVY)
    run_styled(p_inst, "(An Autonomous Institute Affiliated to the University of Mumbai)\n", 'Calibri', 10, italic=True, color=GREY)
    run_styled(p_inst, "Munshi Nagar, Andheri (West), Mumbai — 400 058\n", 'Calibri', 9.5, color=GREY)
    run_styled(p_inst, "NBA Accredited Programs | NAAC 'A' Grade Accredited", 'Calibri', 9.5, bold=True, color=NAVY)

    add_spacer(doc, 48)

    # Magazine Title Card
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = t.cell(0, 0)
    from magazine_core import set_cell_shading, set_cell_margins, set_cell_borders, make_table_full_width
    make_table_full_width(t)
    set_cell_shading(cell, NAVY_HEX)
    set_cell_margins(cell, top=260, bottom=260, left=200, right=200)
    set_cell_borders(cell,
        top={'color': GOLD_HEX, 'sz': '18', 'style': 'single'},
        bottom={'color': GOLD_HEX, 'sz': '18', 'style': 'single'})

    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    run_styled(p_title, "DEPARTMENT OF COMPUTER ENGINEERING\n", 'Calibri', 14, bold=True, color=GOLD)
    run_styled(p_title, "ANNUAL MAGAZINE\n", 'Cambria', 28, bold=True, color=WHITE)
    run_styled(p_title, "ACADEMIC YEAR 2025 – 2026\n", 'Calibri', 14, bold=True, color=WHITE)
    run_styled(p_title, "Innovate • Empower • Transcend", 'Calibri', 11, italic=True, color=GOLD_LIGHT_HEX and RGBColor(0xDC, 0xC8, 0x96))

    add_spacer(doc, 48)

    # Department Key Highlights Banner
    stat_cards(doc, [
        ("98.21%", "Placement Rate"),
        ("₹33.00 LPA", "Highest CTC"),
        ("158", "UG Research Projects"),
        ("40", "Publications (15J + 25C)"),
        ("3,316", "Hackathon Participants")
    ])

    doc.add_page_break()

    # Table of Contents
    section_header(doc, "TABLE OF CONTENTS", "DEPARTMENT OF COMPUTER ENGINEERING — ANNUAL MAGAZINE 2025–26")

    toc_rows = [
        ("1.0", "From the Principal's Desk", "Dr. Bhalchandra Chaudhari"),
        ("2.0", "From the Head of Department's Desk", "Dr. Sudhir N. Dhage"),
        ("3.0", "Vision and Mission Statements", "Institute & Department Strategic Direction"),
        ("4.0", "Program Educational Objectives & Outcomes", "PEOs, POs, PSOs & Audited Attainment"),
        ("5.0", "About the Department & Curriculum Architecture", "Liberal 'Pi-Model' & NEP-2020 Framework"),
        ("6.0", "Laboratory Facilities & Smart Classrooms", "15 Computing Labs & New Equipment Added"),
        ("7.0", "Overview of Academic Programs & Placements", "Statistics, Companies, Internships & Higher Studies"),
        ("8.0", "Faculty and Technical Staff Information", "Faculty Profiles, Designations & Research Domains"),
        ("9.0", "Additional Responsibilities to Faculty", "Institutional Portfolios & Co-curricular Leadership"),
        ("10.0", "Advisory Board, Board of Studies & DAPC", "Curriculum Governance & Academic Quality Monitoring"),
        ("11.0", "Consultancy, Grants, Patents & Publications", "Faculty Citations, Patents & Research Papers"),
        ("12.0", "Faculty and Student Achievements", "Awards, NPTEL Stars, Hackathons & Bug Bounties"),
        ("13.0", "Undergraduate Projects, M.Tech & Ph.D. Research", "158 B.Tech Projects, 15 Dissertations & Doctoral Work"),
        ("14.0", "Industry–Institute Interaction", "Alumni Evaluation, Co-Teaching & Corporate Tracks"),
        ("15.0", "Distinguished Alumni & Mentorship", "Alumni Mentors & Global Professional Footprint"),
        ("16.0", "Departmental Activities & Student Associations", "PTM 2026, SPIT Hackathon, FACE, CSI, IEEE"),
        ("17.0", "Student Support & Innovative Teaching Systems", "Remedial Teaching, Slow/Fast Learners & Pedagogy"),
        ("18.0", "SWOC Analysis & Strategic Roadmap", "Departmental Strengths, Weaknesses, Opportunities & Goals"),
        ("19.0", "Meet the Editorial Board & Acknowledgments", "Student & Faculty Magazine Production Team")
    ]

    data_table(doc, ["Sec.", "Section Title", "Coverage & Highlights"], toc_rows, col_widths=[0.6, 3.8, 2.6])
    doc.add_page_break()

def build_section_01(doc):
    section_header(doc, "From the Principal's Desk", "SECTION 01 — INSTITUTIONAL LEADERSHIP")

    # Layout with Principal Photo
    principal_img = os.path.join(IMG_DIR, "principal_photo.jpg")
    
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    from magazine_core import make_table_full_width, set_cell_margins
    make_table_full_width(t)
    
    cell_img = t.cell(0, 0)
    cell_txt = t.cell(0, 1)
    set_cell_margins(cell_img, top=40, bottom=40, left=40, right=60)
    set_cell_margins(cell_txt, top=40, bottom=40, left=60, right=40)
    
    # Add photo
    if os.path.exists(principal_img):
        p_img = cell_img.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_img.add_run()
        run.add_picture(principal_img, width=Inches(1.8))
        
        p_sub = cell_img.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(4)
        run_styled(p_sub, "Dr. Bhalchandra Chaudhari\n", 'Calibri', 9.5, bold=True, color=NAVY)
        run_styled(p_sub, "Principal\nSardar Patel Institute of Technology", 'Calibri', 8.5, italic=True, color=GREY)

    # Verbatim Text from CSE Magazine AY-2025-2026.pdf (Page 4)
    p1 = cell_txt.paragraphs[0]
    p1.paragraph_format.space_after = Pt(6)
    p1.paragraph_format.line_spacing = 1.15
    run_styled(p1, "I have great pleasure in expressing my thoughts as the Principal of Sardar Patel Institute of Technology (SPIT), the Numero Uno, self-financed, autonomous Institution of Maharashtra. We are a constituent of Bharatiya Vidya Bhavan, not just a conglomeration of more than 300 institutions, but a culture, a saga, a holy journey, started by Dr. K.M. Munshi with the support of Mahatma Gandhi in 1938. Imparting value-based education with Indian cultural ethos has always been the motto of Bhavan.", 'Calibri', 10, color=BLACK)

    p2 = cell_txt.add_paragraph()
    p2.paragraph_format.space_after = Pt(6)
    p2.paragraph_format.line_spacing = 1.15
    run_styled(p2, "Engineers & technologists form the backbone of any nation’s economic development. The world is presently undergoing a very unprecedented, extraordinary, challenging time. A new normal is anticipated in many walks of life, including education. Such disruptions will come again and again in one or the other form. Engineers will work for 40-50 years of their life, they will have 3 to 4 diversified careers in technologies we are even unaware of. To make aspiring minds confident and future-proof, education must prepare them for a “marathon” rather than a “sprint”. SPIT makes continuous, sincere efforts towards this.", 'Calibri', 10, color=BLACK)

    # Remaining paragraphs span full width
    add_spacer(doc, 4)
    body(doc, "We focus on “How to learn?” rather than “What to learn?”. We believe in multidisciplinary exposure to the learners, yet ensuring growth in one vertical, cherishing human sensitivity and empathy. We have thoughtfully articulated a unique academic model towards this. Our splendid academic performance, sparkling placements (quantitative and qualitative), enrollment for higher studies at the best places of the world, prizes won by our students in national/international level technical competitions, in the past many years are the true testimonials for this.")

    body(doc, "With the support of a dedicated and hardworking faculty and staff, the institute has achieved enviable visibility and ranking in a short span. On behalf of all stakeholders of SPIT, I welcome you to this family and look forward to your valuable association with us for a better tomorrow. Four years of engineering education at Sardar Patel Institute of Technology or two years post-graduation, will undoubtedly empower you to lead a successful life.")

    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(6)
    p_end.paragraph_format.space_after = Pt(4)
    run_styled(p_end, "Let’s grow together…\n\n", 'Calibri', 10.5, italic=True, bold=True, color=NAVY)
    run_styled(p_end, "Dr. Bhalchandra Chaudhari\n", 'Cambria', 11.5, bold=True, color=NAVY)
    run_styled(p_end, "Principal,\nBharatiya Vidya Bhavan's Sardar Patel Institute of Technology", 'Calibri', 9.5, italic=True, color=GREY)

    doc.add_page_break()

def build_section_02(doc):
    section_header(doc, "From the Head of Department's Desk", "SECTION 02 — DEPARTMENTAL LEADERSHIP")

    callout_box(doc,
        "\"The year 2025–26 has been a landmark year for the Department of Computer Engineering. Our placement record speaks for itself — with a placement rate of 98.21%, a highest CTC of ₹33 LPA, and an average CTC of ₹10.60 LPA, our students continue to attract the finest technology leaders globally while publishing high-impact research in Scopus and SCIE indexed venues.\"",
        title="Leadership Vision — Dr. Sudhir N. Dhage, Professor & Head",
        border_color=GOLD_HEX, bg_color=GOLD_BG_HEX
    )

    body(doc, "It is with great pride and satisfaction that I present the Annual Magazine of the Department of Computer Engineering for the academic year 2025–26. This report encapsulates the achievements, progress, and aspirations of our department across academics, research, placements, and co-curricular engagement.")

    body(doc, "The year 2025–26 has been a landmark year for the department. Our placement record speaks for itself — with a placement percentage of 98.21%, a highest CTC of ₹33 LPA, and an average CTC of ₹10.60 LPA, our students continue to attract the finest companies globally, including Barclays, Nomura, BlackRock, JP Morgan, PhonePe, Morgan Stanley, and many more. Equally heartening is the fact that 32 students have secured admissions to premier institutions including IIT Bombay, IIT Madras, IIT Patna, SPJIMR, and DRDO for higher studies.")

    body(doc, "On the research front, the department has produced 15 international journal publications and 25 conference papers, with 9 Scopus-indexed and 3 SCIE-indexed publications. We have also filed a patent for an AI-powered interactive voice-enabled drawing tool — a testament to the innovation culture we foster.")

    body(doc, "Our students have brought laurels to the department at national and international competitions — winning the EY Techathon 6.0, qualifying as ICPC Asia Amritapuri Regionalists, earning spots on the Ethereum and Zoho Bug Bounty Halls of Fame, and winning at IIT Bombay, IIM Indore, and Google Cloud hackathons. The 10th edition of the SPIT Hackathon attracted 3,316 registrations from across the country.")

    body(doc, "I extend my heartfelt gratitude to our visionary Principal, Dr. Bhalchandra Chaudhari, the management of Bharatiya Vidya Bhavan, our dedicated faculty members, technical staff, industry partners, alumni, parents, and most importantly, our students, who are the cornerstone of everything we achieve.")

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(8)
    p_sig.paragraph_format.space_after = Pt(4)
    run_styled(p_sig, "Dr. Sudhir N. Dhage\n", 'Cambria', 11.5, bold=True, color=NAVY)
    run_styled(p_sig, "Professor & Head, Department of Computer Engineering\n", 'Calibri', 10, bold=True, color=GOLD_DARK)
    run_styled(p_sig, "Bharatiya Vidya Bhavan's Sardar Patel Institute of Technology", 'Calibri', 9.5, italic=True, color=GREY)

    doc.add_page_break()

def build_section_03_and_04(doc):
    section_header(doc, "Vision, Mission, PEOs, POs, PSOs & Attainment", "SECTIONS 03 & 04 — STRATEGIC ACADEMIC FRAMEWORK")

    sub_section_header(doc, "Institute Strategic Foundation", "3.1")
    callout_box(doc,
        "\"To be a leading technical institution constantly pursuing excellence to develop holistic contributors for the society.\"",
        title="INSTITUTE VISION (Audit PDF Page 4)",
        border_color=NAVY_HEX, bg_color=LIGHT_BG
    )

    body(doc, "Institute Mission:", bold=True, color=NAVY)
    bullet(doc, "To provide a conducive academic and research environment that stimulates learning, promotes intellectual growth, and ensures professional development of students, faculty, and staff.")
    bullet(doc, "To collaborate with renowned industry, academic institutions, and research laboratories for academics and research to solve real life problems.")
    bullet(doc, "To leverage the strength of alumni to foster industry-relevant coaching and mentoring of the students.")
    bullet(doc, "To nurture values with a sense of social responsibility and inculcate attitude and skills required for holistic development of the students.")

    add_spacer(doc, 4)

    sub_section_header(doc, "Department of Computer Engineering Strategic Foundation", "3.2")
    callout_box(doc,
        "\"To build a strong teaching and research environment in the domain of Computer Engineering to foster innovative, entrepreneurial and socially responsible contributors.\"",
        title="DEPARTMENT VISION (Audit PDF Page 5)",
        border_color=GOLD_HEX, bg_color=GOLD_BG_HEX
    )

    body(doc, "Department Mission:", bold=True, color=NAVY)
    bullet(doc, "To provide state-of-the-art infrastructure and curriculum to enhance the knowledge, skills and capabilities of students.")
    bullet(doc, "To develop an ecosystem conducive to research and development through industry-academia collaboration.")
    bullet(doc, "To inculcate ethical values, professional discipline, leadership qualities, and social responsiveness among students.")
    bullet(doc, "To foster entrepreneurship and innovative problem solving to address real-world socio-economic challenges.")

    add_spacer(doc, 6)

    sub_section_header(doc, "Program Educational Objectives (PEOs) & Program Specific Outcomes (PSOs)", "4.1")
    body(doc, "B.Tech. Computer Engineering — PEOs (Audit PDF Page 66):", bold=True, color=NAVY)
    bullet(doc, "Graduates will demonstrate proficiency in computing technologies and innovate in design, development, deployment, and maintenance of serving various sectors of economy.", bold_prefix="PEO1:")
    bullet(doc, "Graduates will be effective communicators, inspiring leaders, aspiring team members, and trailblazers of computational technologies.", bold_prefix="PEO2:")
    bullet(doc, "Graduates will be contributory professionals demonstrating lifelong learning attitude and sensitivity to society problems.", bold_prefix="PEO3:")

    body(doc, "B.Tech. Computer Engineering — PSOs (Audit PDF Page 66):", bold=True, color=NAVY)
    bullet(doc, "Possess strong foundation of Computer Engineering (Algorithms, Data Structures, AIML, Databases, Networks, Security etc.) and apply the same for finding efficient and feasible solutions for real-life problems.", bold_prefix="PSO1:")
    bullet(doc, "Be motivated to engage in lifelong learning, adopt strong values, and develop skills in project and finance management. Showcase leadership, teamwork, and social awareness, while becoming proficient in databases, application software, computer networks and security, cloud services, storage, data manipulation, etc.", bold_prefix="PSO2:")

    add_spacer(doc, 4)

    body(doc, "M.Tech. Computer Engineering — POs & PEOs (Audit PDF Page 67):", bold=True, color=NAVY)
    bullet(doc, "An ability to independently carry out research / investigation and development work to solve practical problems.", bold_prefix="PO1:")
    bullet(doc, "An ability to write and present a substantial technical report/document.", bold_prefix="PO2:")
    bullet(doc, "Students should be able to demonstrate a degree of mastery over the area as per the specialization of the program. The mastery should be at a level higher than the requirements in the appropriate bachelor program.", bold_prefix="PO3:")
    bullet(doc, "Graduates will demonstrate technical competency and maintain software products, services and systems with strong analytical and problem-solving skills.", bold_prefix="PEO1:")
    bullet(doc, "Graduates will be adept in collaborating on interdisciplinary projects fostering innovations.", bold_prefix="PEO2:")
    bullet(doc, "Graduates will exhibit effective leadership, teamwork, and lifelong learning.", bold_prefix="PEO3:")

    add_spacer(doc, 6)

    sub_section_header(doc, "Audited Program Outcomes & PSOs Attainment Matrix", "4.2")
    body(doc, "As verified during the External Academic Audit (Audit PDF Page 37), the department computes course attainment every semester, PO/PSO attainment annually, and PEO attainment at four-year intervals using calibrated direct and indirect assessment tools.")

    # Attainment Table 1: Batch 2021-25
    body(doc, "Batch 2021–25: PO & PSO Attainment (Audited Values)", bold=True, color=NAVY)
    b1_headers = ["Metric", "PO1", "PO2", "PO3", "PO4", "PO5", "PO6", "PO7", "PO8", "PO9", "PO10", "PO11", "PO12", "PSO1", "PSO2"]
    b1_rows = [
        ["Target", "2.29", "2.21", "2.33", "1.85", "2.45", "1.86", "1.13", "2.07", "2.02", "2.13", "1.83", "1.82", "2.15", "1.72"],
        ["Direct", "1.75", "1.76", "1.73", "1.67", "1.90", "1.49", "1.35", "1.22", "1.59", "1.69", "1.79", "1.49", "1.39", "1.16"],
        ["Indirect", "2.40", "2.48", "2.37", "2.22", "2.41", "1.76", "1.06", "2.23", "2.33", "2.58", "2.37", "2.44", "1.87", "2.26"],
        ["Attainment", "1.88", "1.91", "1.86", "1.78", "2.00", "1.54", "1.29", "1.42", "1.74", "1.87", "1.90", "1.68", "1.49", "1.38"]
    ]
    data_table(doc, b1_headers, b1_rows, compact=True)

    # Attainment Table 2: Batch 2022-26
    body(doc, "Batch 2022–26: PO & PSO Attainment (Audited Values)", bold=True, color=NAVY)
    b2_headers = ["Metric", "PO1", "PO2", "PO3", "PO4", "PO5", "PO6", "PO7", "PO8", "PO9", "PO10", "PO11", "PO12", "PSO1", "PSO2"]
    b2_rows = [
        ["Target", "2.39", "2.14", "2.19", "1.74", "2.29", "2.20", "1.40", "2.36", "2.25", "2.10", "1.83", "1.93", "2.00", "1.80"],
        ["Direct", "1.88", "1.66", "1.66", "1.41", "1.79", "1.59", "1.10", "1.52", "1.54", "1.41", "1.47", "1.52", "1.17", "1.29"],
        ["Indirect", "2.44", "2.25", "2.21", "1.73", "2.25", "1.80", "1.05", "2.51", "2.32", "2.19", "1.56", "2.08", "1.33", "1.45"],
        ["Attainment", "1.99", "1.78", "1.77", "1.47", "1.88", "1.63", "1.09", "1.72", "1.69", "1.56", "1.49", "1.63", "1.20", "1.32"]
    ]
    data_table(doc, b2_headers, b2_rows, compact=True)

    doc.add_page_break()
